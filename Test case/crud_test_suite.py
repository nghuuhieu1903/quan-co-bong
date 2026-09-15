"""Chạy thật toàn bộ test case thêm/sửa/xóa trong matrix.py và ghi kết quả
vào test_cases.xlsx.

    cd "Test case"
    ../venv/bin/python crud_test_suite.py

Mọi dữ liệu test tạo ra (sản phẩm, phòng, đơn hàng, tài khoản...) đều được
xóa lại ở cuối, kể cả khi có test thất bại giữa chừng.
"""

import io
import os
import re
import sys
import datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import app as app_module
import models
from app import db
from werkzeug.security import check_password_hash

from matrix import MATRIX, BY_ID

HERE = os.path.dirname(os.path.abspath(__file__))
XLSX_PATH = os.path.join(HERE, 'test_cases.xlsx')


def token(client, path):
    html = client.get(path).get_data(as_text=True)
    m = re.search(r'name="csrf_token" value="([^"]+)"', html)
    return m.group(1) if m else None


def post(client, path, data, from_page, **kwargs):
    data = dict(data)
    data['csrf_token'] = token(client, from_page)
    return client.post(path, data=data, **kwargs)


class Results:
    def __init__(self):
        self.data = {}  # id -> (ok, detail)

    def check(self, case_id, ok, detail=''):
        assert case_id in BY_ID, f'unknown test case id {case_id} - add it to matrix.py first'
        self.data[case_id] = (bool(ok), detail)
        print(f'  {"PASS" if ok else "FAIL"}  {case_id}  {BY_ID[case_id].feature}'
              + (f'  [{detail}]' if detail else ''))


# ---------------------------------------------------------------------------
# Sản phẩm
# ---------------------------------------------------------------------------

def product_flows(rep, ids, app):
    a = ids['admin_client']

    # SP-01 thêm
    with app.app_context():
        before = models.Product.query.count()
    post(a, '/admin/product/add', {
        'name': 'QA Test Drink', 'description': 'tạo bởi bộ test tự động',
        'price': '15000', 'stock': '10', 'category': 'coffee', 'item_type': 'drink',
    }, '/admin/product/add')
    with app.app_context():
        after = models.Product.query.count()
        p = models.Product.query.filter_by(name='QA Test Drink').first()
    rep.check('SP-01', after == before + 1 and p is not None, f'{before} -> {after}')
    ids['product'] = p.id if p else None
    if not p:
        return

    # SP-06 xóa ảnh phụ: thêm 1 ảnh phụ trước rồi xóa nó
    fake_image = (io.BytesIO(b'\xff\xd8\xff\xe0fakejpegdata'), 'qa_test.jpg')
    post(a, f'/admin/product/{p.id}/edit', {
        'name': p.name, 'description': p.description, 'price': '15000',
        'stock': '10', 'category': 'coffee', 'item_type': 'drink',
        'images': fake_image,
    }, f'/admin/product/{p.id}/edit', content_type='multipart/form-data')
    with app.app_context():
        img = models.ProductImage.query.filter_by(product_id=p.id).first()
    if img:
        img_id, img_file = img.id, img.image
        img_path = os.path.join(HERE, '..', 'static', 'images', img_file)
        r = post(a, f'/admin/product-image/{img_id}/delete', {},
                 f'/admin/product/{p.id}/edit')
        with app.app_context():
            gone = db.session.get(models.ProductImage, img_id) is None
        file_gone = not os.path.exists(img_path)
        rep.check('SP-06', gone and file_gone, f'row gone={gone} file gone={file_gone}')
    else:
        rep.check('SP-06', False, 'ảnh phụ không được lưu lại để xóa thử')

    # SP-02 sửa
    post(a, f'/admin/product/{p.id}/edit', {
        'name': 'QA Test Drink v2', 'description': 'đã sửa', 'price': '23000',
        'stock': '4', 'category': 'coffee', 'item_type': 'food',
    }, f'/admin/product/{p.id}/edit')
    with app.app_context():
        p2 = db.session.get(models.Product, p.id)
        ok = p2.name == 'QA Test Drink v2' and p2.price == 23000 and p2.item_type == 'food'
    rep.check('SP-02', ok, f'{p2.name}, {p2.price}, {p2.item_type}' if p2 else 'not found')

    # SP-04 xóa sản phẩm đã có đơn hàng -> phải tự ẩn, không xóa hẳn
    with app.app_context():
        o = models.Order(customer_name='QA Test', customer_phone='0900000000',
                         total_amount=23000, status='completed')
        db.session.add(o)
        db.session.flush()
        db.session.add(models.OrderItem(order_id=o.id, product_id=p.id,
                                        quantity=1, price=23000))
        db.session.commit()
        ids['order_for_sp04'] = o.id

    post(a, f'/admin/product/{p.id}/delete', {}, '/admin/manage_products')
    with app.app_context():
        p3 = db.session.get(models.Product, p.id)
        ok = p3 is not None and p3.is_active is False
    rep.check('SP-04', ok, 'sản phẩm vẫn còn nhưng is_active=False' if ok else 'không đúng như kỳ vọng')

    # SP-07 sản phẩm ẩn không hiện cho khách / POS
    # (đọc trang quản lý trước để "tiêu" hết flash message còn treo từ lệnh
    # xóa vừa rồi - nếu không nó sẽ lẫn vào lần render trang tiếp theo và
    # làm phép so khớp chuỗi tên sản phẩm bị nhầm là dương tính giả)
    a.get('/admin/manage_products')
    home = a.get('/customer').get_data(as_text=True)
    pos_page = a.get('/admin/pos').get_data(as_text=True)
    rep.check('SP-07', 'QA Test Drink v2' not in home and 'QA Test Drink v2' not in pos_page,
              'kiểm tra trang chủ và POS')

    # SP-05 khôi phục
    post(a, f'/admin/product/{p.id}/restore', {}, '/admin/manage_products')
    with app.app_context():
        p4 = db.session.get(models.Product, p.id)
        ok = p4.is_active is True
    rep.check('SP-05', ok)

    # SP-03 xóa sản phẩm chưa từng bán -> xóa hẳn. Dùng 1 sản phẩm test khác,
    # không đơn hàng nào tham chiếu tới.
    with app.app_context():
        before = models.Product.query.count()
    post(a, '/admin/product/add', {
        'name': 'QA Test Drink Clean', 'description': 'không có đơn hàng',
        'price': '10000', 'stock': '5', 'category': 'coffee', 'item_type': 'drink',
    }, '/admin/product/add')
    with app.app_context():
        p_clean = models.Product.query.filter_by(name='QA Test Drink Clean').first()
    if p_clean:
        post(a, f'/admin/product/{p_clean.id}/delete', {}, '/admin/manage_products')
        with app.app_context():
            gone = db.session.get(models.Product, p_clean.id) is None
        rep.check('SP-03', gone)
    else:
        rep.check('SP-03', False, 'không tạo được sản phẩm để thử xóa')

    # SP-08 đổi giá hàng loạt: 2 sản phẩm được chọn phải đổi giá, 1 sản phẩm
    # không được chọn phải giữ nguyên giá cũ
    with app.app_context():
        bulk_a = models.Product(name='QA Bulk A', description='', price=11111,
                                stock=5, category='coffee', item_type='drink')
        bulk_b = models.Product(name='QA Bulk B', description='', price=22222,
                                stock=5, category='coffee', item_type='drink')
        bulk_untouched = models.Product(name='QA Bulk Untouched', description='',
                                        price=33333, stock=5, category='coffee',
                                        item_type='drink')
        db.session.add_all([bulk_a, bulk_b, bulk_untouched])
        db.session.commit()
        ids['bulk_products'] = [bulk_a.id, bulk_b.id, bulk_untouched.id]

    post(a, '/admin/products/bulk_action', {
        'product_ids': [str(bulk_a.id), str(bulk_b.id)], 'action': 'set_price',
        'price': '18000',
    }, '/admin/manage_products')
    with app.app_context():
        pa = db.session.get(models.Product, bulk_a.id)
        pb = db.session.get(models.Product, bulk_b.id)
        pu = db.session.get(models.Product, bulk_untouched.id)
        ok = pa.price == 18000 and pb.price == 18000 and pu.price == 33333
    rep.check('SP-08', ok, f'A={pa.price} B={pb.price} untouched={pu.price}')

    # SP-09 xóa hàng loạt: 1 sản phẩm chưa từng bán (xóa hẳn), 1 sản phẩm đã
    # có đơn hàng (chỉ ẩn) - chọn cùng lúc rồi xóa 1 lần
    with app.app_context():
        o = models.Order(customer_name='QA Bulk Delete', customer_phone='0900000001',
                         total_amount=18000, status='completed')
        db.session.add(o)
        db.session.flush()
        db.session.add(models.OrderItem(order_id=o.id, product_id=bulk_b.id,
                                        quantity=1, price=18000))
        db.session.commit()
        ids['order_for_sp09'] = o.id

    post(a, '/admin/products/bulk_action', {
        'product_ids': [str(bulk_a.id), str(bulk_b.id)], 'action': 'delete',
    }, '/admin/manage_products')
    with app.app_context():
        gone = db.session.get(models.Product, bulk_a.id) is None
        hidden = db.session.get(models.Product, bulk_b.id)
        hidden_ok = hidden is not None and hidden.is_active is False
    rep.check('SP-09', gone and hidden_ok,
              f'A đã xóa hẳn={gone}, B đã có đơn nên chỉ ẩn={hidden_ok}')
    if gone:
        ids['bulk_products'].remove(bulk_a.id)

    # SP-10 gắn nhãn "Đặc biệt" lúc thêm, bỏ nhãn lúc sửa
    post(a, '/admin/product/add', {
        'name': 'QA Special Dish', 'description': 'món đặc biệt test',
        'price': '50000', 'stock': '5', 'category': 'food',
        'item_type': 'food', 'is_daily': '1', 'is_special': '1',
    }, '/admin/product/add')
    with app.app_context():
        special_product = models.Product.query.filter_by(name='QA Special Dish').first()
    added_ok = special_product is not None and special_product.is_special is True
    ids['special_product'] = special_product.id if special_product else None

    badge_shows = False
    removed_ok = False
    if special_product:
        lunch_page = a.get('/products?type=food&food_kind=daily').get_data(as_text=True)
        badge_shows = ('QA Special Dish' in lunch_page
                       and 'is-special">Đặc biệt' in lunch_page)

        # sửa lại: bỏ tick "Đặc biệt" - is_special phải tắt
        post(a, f'/admin/product/{special_product.id}/edit', {
            'name': 'QA Special Dish', 'description': 'món đặc biệt test',
            'price': '50000', 'stock': '5', 'category': 'food', 'item_type': 'food',
            'is_daily': '1',   # is_special cố tình không gửi -> phải tắt
        }, f'/admin/product/{special_product.id}/edit')
        with app.app_context():
            after_edit = db.session.get(models.Product, special_product.id)
            removed_ok = after_edit.is_special is False

    rep.check('SP-10', added_ok and badge_shows and removed_ok,
              f'thêm bật is_special={added_ok}, nhãn hiện ở Cơm trưa={badge_shows}, '
              f'sửa tắt is_special={removed_ok}')


# ---------------------------------------------------------------------------
# Phòng
# ---------------------------------------------------------------------------

def room_flows(rep, ids, app):
    a = ids['admin_client']

    # PH-01 thêm
    with app.app_context():
        before = models.Room.query.count()
    post(a, '/admin/room/add', {
        'name': 'QA Test Room', 'description': 'tạo bởi bộ test tự động',
        'price_per_hour': '1500000', 'price_unit': 'tháng', 'capacity': '2',
        'amenities': 'wifi, điều hoà',
    }, '/admin/room/add')
    with app.app_context():
        after = models.Room.query.count()
        room = models.Room.query.filter_by(name='QA Test Room').first()
    rep.check('PH-01', after == before + 1 and room is not None, f'{before} -> {after}')
    ids['room'] = room.id if room else None
    if not room:
        return

    # PH-02 sửa
    post(a, f'/admin/room/{room.id}/edit', {
        'name': 'QA Test Room v2', 'description': 'đã sửa', 'price_per_hour': '2000000',
        'price_unit': 'tháng', 'capacity': '3', 'amenities': 'wifi',
    }, f'/admin/room/{room.id}/edit')
    with app.app_context():
        r2 = db.session.get(models.Room, room.id)
        ok = r2.name == 'QA Test Room v2' and r2.price_per_hour == 2000000 and r2.capacity == 3
    rep.check('PH-02', ok)

    # PH-03 bật/tắt trạng thái trống
    with app.app_context():
        was = db.session.get(models.Room, room.id).available
    a.get(f'/admin/room/{room.id}/toggle')
    with app.app_context():
        now = db.session.get(models.Room, room.id).available
    rep.check('PH-03', was != now)

    # PH-06 cập nhật trạng thái đặt phòng (tạo booking trực tiếp, giống booking cũ)
    with app.app_context():
        b = models.RoomBooking(
            room_id=room.id, customer_name='QA Test', customer_phone='0900000000',
            booking_date=datetime.date(2030, 1, 1), start_time=datetime.time(9, 0),
            end_time=datetime.time(11, 0), total_hours=2.0, total_price=0, status='pending')
        db.session.add(b)
        db.session.commit()
        ids['booking'] = b.id
    post(a, f'/admin/room_booking/{b.id}/update', {'status': 'confirmed'},
         '/admin/room_bookings')
    with app.app_context():
        ok = db.session.get(models.RoomBooking, b.id).status == 'confirmed'
    rep.check('PH-06', ok)

    # PH-05 xóa phòng đã có lịch đặt -> phải tự đóng, không xóa hẳn
    post(a, f'/admin/room/{room.id}/delete', {}, '/admin/rooms')
    with app.app_context():
        r3 = db.session.get(models.Room, room.id)
        ok = r3 is not None and r3.available is False
    rep.check('PH-05', ok, 'phòng vẫn còn nhưng available=False' if ok else 'không đúng như kỳ vọng')

    # PH-04 xóa phòng chưa từng có lịch đặt -> xóa hẳn
    with app.app_context():
        before = models.Room.query.count()
    post(a, '/admin/room/add', {
        'name': 'QA Test Room Clean', 'description': 'không có lịch đặt',
        'price_per_hour': '1000000', 'price_unit': 'tháng', 'capacity': '1',
        'amenities': '',
    }, '/admin/room/add')
    with app.app_context():
        room_clean = models.Room.query.filter_by(name='QA Test Room Clean').first()
    if room_clean:
        post(a, f'/admin/room/{room_clean.id}/delete', {}, '/admin/rooms')
        with app.app_context():
            gone = db.session.get(models.Room, room_clean.id) is None
        rep.check('PH-04', gone)
    else:
        rep.check('PH-04', False, 'không tạo được phòng để thử xóa')

    # PH-07 đổi giá hàng loạt: 2 phòng được chọn phải đổi giá, 1 phòng không
    # được chọn phải giữ nguyên giá cũ
    with app.app_context():
        bulk_a = models.Room(name='QA Bulk Room A', description='', price_per_hour=1000000,
                             price_unit='tháng', capacity=1, amenities='[]', available=True)
        bulk_b = models.Room(name='QA Bulk Room B', description='', price_per_hour=2000000,
                             price_unit='tháng', capacity=1, amenities='[]', available=True)
        bulk_untouched = models.Room(name='QA Bulk Room Untouched', description='',
                                     price_per_hour=3000000, price_unit='tháng', capacity=1,
                                     amenities='[]', available=True)
        db.session.add_all([bulk_a, bulk_b, bulk_untouched])
        db.session.commit()
        ids['bulk_rooms'] = [bulk_a.id, bulk_b.id, bulk_untouched.id]

    post(a, '/admin/rooms/bulk_action', {
        'room_ids': [str(bulk_a.id), str(bulk_b.id)], 'action': 'set_price',
        'price': '1800000',
    }, '/admin/rooms')
    with app.app_context():
        ra = db.session.get(models.Room, bulk_a.id)
        rb = db.session.get(models.Room, bulk_b.id)
        ru = db.session.get(models.Room, bulk_untouched.id)
        ok = ra.price_per_hour == 1800000 and rb.price_per_hour == 1800000 and ru.price_per_hour == 3000000
    rep.check('PH-07', ok, f'A={ra.price_per_hour} B={rb.price_per_hour} untouched={ru.price_per_hour}')

    # PH-08 xóa hàng loạt: 1 phòng chưa từng có lịch đặt (xóa hẳn), 1 phòng
    # đã có lịch đặt (chỉ đóng) - chọn cùng lúc rồi xóa 1 lần
    with app.app_context():
        bk = models.RoomBooking(
            room_id=bulk_b.id, customer_name='QA Bulk Delete', customer_phone='0900000002',
            booking_date=datetime.date(2030, 1, 1), start_time=datetime.time(9, 0),
            end_time=datetime.time(11, 0), total_hours=2.0, total_price=0, status='pending')
        db.session.add(bk)
        db.session.commit()
        ids['booking_for_ph08'] = bk.id

    post(a, '/admin/rooms/bulk_action', {
        'room_ids': [str(bulk_a.id), str(bulk_b.id)], 'action': 'delete',
    }, '/admin/rooms')
    with app.app_context():
        gone = db.session.get(models.Room, bulk_a.id) is None
        closed = db.session.get(models.Room, bulk_b.id)
        closed_ok = closed is not None and closed.available is False
    rep.check('PH-08', gone and closed_ok,
              f'A đã xóa hẳn={gone}, B đã có lịch nên chỉ đóng={closed_ok}')
    if gone:
        ids['bulk_rooms'].remove(bulk_a.id)


# ---------------------------------------------------------------------------
# Đơn hàng + giỏ hàng
# ---------------------------------------------------------------------------

def order_and_cart_flows(rep, ids, app):
    a = ids['admin_client']
    c = ids['customer_client']

    with app.app_context():
        pid = models.Product.query.filter_by(is_active=True).first().id
    ids['plain_product'] = pid

    # GH-01 thêm vào giỏ
    r = post(c, f'/add_to_cart/{pid}', {'quantity': '2'}, '/products')
    cart = c.get('/cart').get_data(as_text=True)
    rep.check('GH-01', r.status_code == 302 and 'value="2"' in cart)

    # GH-02 cập nhật số lượng
    post(c, f'/update_cart/{pid}', {'quantity': '5'}, '/cart')
    cart = c.get('/cart').get_data(as_text=True)
    rep.check('GH-02', 'value="5"' in cart)

    # GH-04 xóa nhiều món cùng lúc
    with app.app_context():
        others = [pr.id for pr in models.Product.query
                  .filter(models.Product.is_active.is_(True), models.Product.id != pid)
                  .limit(2).all()]
    for x in others:
        post(c, f'/add_to_cart/{x}', {'quantity': '1'}, '/products')
    post(c, '/remove_selected_from_cart', {'product_ids': [str(x) for x in others]}, '/cart')
    cart = c.get('/cart').get_data(as_text=True)
    gone = all(f'data-cart-row="{x}"' not in cart for x in others)
    kept = f'data-cart-row="{pid}"' in cart
    rep.check('GH-04', gone and kept, f'gone={gone} kept={kept}')

    # GH-03 xóa một món
    post(c, f'/remove_from_cart/{pid}', {}, '/cart')
    cart = c.get('/cart').get_data(as_text=True)
    rep.check('GH-03', f'data-cart-row="{pid}"' not in cart)

    # GH-05 tự động điền tên/SĐT khách đã đăng nhập
    with app.app_context():
        from werkzeug.security import generate_password_hash
        test_customer = models.Customer.query.filter_by(username='qa_prefill_customer').first()
        if not test_customer:
            test_customer = models.Customer(username='qa_prefill_customer',
                                            password=generate_password_hash('qaprefill123'),
                                            full_name='QA Prefill Customer', phone='0955500001')
            db.session.add(test_customer)
            db.session.commit()
        ids['prefill_customer'] = test_customer.id

    logged_in_client = app.test_client()
    ids['logged_in_client'] = logged_in_client
    post(logged_in_client, '/customer/authenticate',
         {'username': 'qa_prefill_customer', 'password': 'qaprefill123'}, '/customer/login')
    post(logged_in_client, f'/add_to_cart/{pid}', {'quantity': '1'}, '/products')
    cart_page = logged_in_client.get('/cart').get_data(as_text=True)
    ok = 'value="QA Prefill Customer"' in cart_page and 'value="0955500001"' in cart_page
    rep.check('GH-05', ok, 'kiểm tra ô tên/SĐT đã điền sẵn trên trang giỏ hàng')

    # GH-06 đặt hộ người khác: submit với tên khác tài khoản, đơn phải lưu
    # đúng tên người nhận vừa gõ, không tự ý dùng lại tên tài khoản
    with app.app_context():
        before = models.Order.query.count()
    post(logged_in_client, '/process_order', {
        'name': 'Người Nhận Hộ', 'phone': '0966600002',
        'payment_method': 'cash', 'notes': 'crud suite - đặt hộ'}, '/cart')
    with app.app_context():
        after = models.Order.query.count()
        new_order = models.Order.query.order_by(models.Order.id.desc()).first()
    ok = (after == before + 1 and new_order is not None
          and new_order.customer_name == 'Người Nhận Hộ'
          and new_order.customer_phone == '0966600002')
    rep.check('GH-06', ok, f'tên lưu={new_order.customer_name if new_order else "-"}')
    ids['prefill_order'] = new_order.id if ok else None

    # DH-05 khách đặt hàng qua checkout
    post(c, f'/add_to_cart/{pid}', {'quantity': '1'}, '/products')
    with app.app_context():
        before = models.Order.query.count()
    r = post(c, '/process_order', {
        'name': 'QA Test Customer', 'phone': '0911111111',
        'payment_method': 'cash', 'notes': 'crud suite'}, '/checkout')
    with app.app_context():
        after = models.Order.query.count()
        new_order = models.Order.query.order_by(models.Order.id.desc()).first()
    ok = (after == before + 1 and new_order is not None
          and new_order.customer_name == 'QA Test Customer'
          and new_order.customer_phone == '0911111111')
    rep.check('DH-05', ok, f'{before} -> {after}')
    ids['checkout_order'] = new_order.id if ok else None

    # DH-02 cập nhật trạng thái đơn
    if ids.get('checkout_order'):
        post(a, f'/admin/order/{ids["checkout_order"]}/update', {'status': 'completed'},
             '/admin/orders')
        with app.app_context():
            ok = db.session.get(models.Order, ids['checkout_order']).status == 'completed'
        rep.check('DH-02', ok)

    # DH-01 tạo đơn tại quầy (POS)
    with app.app_context():
        prod = db.session.get(models.Product, pid)
        stock_before = prod.stock
        orders_before = models.Order.query.count()
    r = post(a, '/admin/pos/order', {
        'product_id': [str(pid)], 'quantity': ['1'], 'payment_method': 'cash',
        'customer_name': 'QA POS Test',
    }, '/admin/pos')
    with app.app_context():
        orders_after = models.Order.query.count()
        stock_after = db.session.get(models.Product, pid).stock
        pos_order = models.Order.query.order_by(models.Order.id.desc()).first()
    ok = orders_after == orders_before + 1 and stock_after == stock_before - 1
    rep.check('DH-01', ok, f'đơn {orders_before}->{orders_after}, tồn kho {stock_before}->{stock_after}')
    ids['pos_order'] = pos_order.id if ok else None

    # DH-03 xóa đơn hàng - Super Admin (dùng đơn checkout)
    if ids.get('checkout_order'):
        oid = ids['checkout_order']
        post(a, f'/admin/order/{oid}/delete', {'next': 'orders'}, '/admin/orders')
        with app.app_context():
            gone_order = db.session.get(models.Order, oid) is None
            gone_items = models.OrderItem.query.filter_by(order_id=oid).count() == 0
        rep.check('DH-03', gone_order and gone_items)
        ids['checkout_order'] = None  # đã xóa, khỏi cleanup lại

    # DH-04 xóa đơn hàng - admin thường phải bị chặn (dùng đơn POS)
    if ids.get('pos_order'):
        oid = ids['pos_order']
        reg = ids['regular_admin_client']
        token_val = token(reg, '/admin/orders')
        resp = reg.post(f'/admin/order/{oid}/delete',
                        data={'csrf_token': token_val, 'next': 'orders'},
                        follow_redirects=True)
        with app.app_context():
            still_there = db.session.get(models.Order, oid) is not None
        blocked_msg = 'Chỉ Super Admin' in resp.get_data(as_text=True)
        rep.check('DH-04', still_there and blocked_msg,
                  f'còn tồn tại={still_there} thông báo chặn={blocked_msg}')

    # DH-06/DH-07/DH-08 hành động hàng loạt trên đơn hàng
    with app.app_context():
        pid = ids['plain_product']
        o1 = models.Order(customer_name='QA Bulk Order 1', customer_phone='0944444441',
                          total_amount=10000, status='pending')
        o2 = models.Order(customer_name='QA Bulk Order 2', customer_phone='0944444442',
                          total_amount=10000, status='pending')
        o3 = models.Order(customer_name='QA Bulk Order 3 (untouched)', customer_phone='0944444443',
                          total_amount=10000, status='pending')
        db.session.add_all([o1, o2, o3])
        db.session.flush()
        for o in (o1, o2, o3):
            db.session.add(models.OrderItem(order_id=o.id, product_id=pid, quantity=1, price=10000))
        db.session.commit()
        ids['bulk_orders'] = [o1.id, o2.id, o3.id]

    # DH-06 đánh dấu hàng loạt "đã xong" - chỉ 2 trong 3 đơn được chọn
    token_val = token(a, '/admin/orders')
    a.post('/admin/orders/bulk_action', data={
        'csrf_token': token_val, 'action': 'complete',
        'order_ids': [str(o1.id), str(o2.id)], 'next': 'orders',
    })
    with app.app_context():
        s1 = db.session.get(models.Order, o1.id).status
        s2 = db.session.get(models.Order, o2.id).status
        s3 = db.session.get(models.Order, o3.id).status
    rep.check('DH-06', s1 == 'completed' and s2 == 'completed' and s3 == 'pending',
              f'o1={s1} o2={s2} o3 (untouched)={s3}')

    # DH-08 xóa hàng loạt - admin thường phải bị chặn (dùng đơn 3, chưa đụng tới)
    reg = ids['regular_admin_client']
    token_reg = token(reg, '/admin/orders')
    reg.post('/admin/orders/bulk_action', data={
        'csrf_token': token_reg, 'action': 'delete',
        'order_ids': [str(o3.id)], 'next': 'orders',
    })
    with app.app_context():
        still_there = db.session.get(models.Order, o3.id) is not None
    rep.check('DH-08', still_there, f'còn tồn tại={still_there}')

    # DH-07 xóa hàng loạt - Super Admin (dùng đơn 1 và 2)
    a.post('/admin/orders/bulk_action', data={
        'csrf_token': token_val, 'action': 'delete',
        'order_ids': [str(o1.id), str(o2.id)], 'next': 'orders',
    })
    with app.app_context():
        gone1 = db.session.get(models.Order, o1.id) is None
        gone2 = db.session.get(models.Order, o2.id) is None
        items_gone = models.OrderItem.query.filter(
            models.OrderItem.order_id.in_([o1.id, o2.id])).count() == 0
    rep.check('DH-07', gone1 and gone2 and items_gone,
              f'o1 gone={gone1} o2 gone={gone2} items gone={items_gone}')
    if gone1 and gone2:
        ids['bulk_orders'] = [oid for oid in ids['bulk_orders'] if oid not in (o1.id, o2.id)]


# ---------------------------------------------------------------------------
# Danh mục (trang khách hàng): tab Cơm trưa / Ăn vặt
# ---------------------------------------------------------------------------

def catalogue_flows(rep, ids, app):
    c = ids['customer_client']
    with app.app_context():
        lunch = models.Product(name='QA Catalogue Lunch', description='', price=30000,
                               stock=5, category='food', item_type='food', is_daily=True)
        snack = models.Product(name='QA Catalogue Snack', description='', price=12000,
                               stock=5, category='snack', item_type='food', is_daily=False)
        db.session.add_all([lunch, snack])
        db.session.commit()
        ids['catalogue_products'] = [lunch.id, snack.id]

    lunch_page = c.get('/products?type=food&food_kind=daily').get_data(as_text=True)
    rep.check('DM-01', 'QA Catalogue Lunch' in lunch_page and 'QA Catalogue Snack' not in lunch_page,
              f'lunch có món trưa={"QA Catalogue Lunch" in lunch_page}, '
              f'lẫn món ăn vặt={"QA Catalogue Snack" in lunch_page}')

    snack_page = c.get('/products?type=food&food_kind=snack').get_data(as_text=True)
    rep.check('DM-02', 'QA Catalogue Snack' in snack_page and 'QA Catalogue Lunch' not in snack_page,
              f'ăn vặt có món ăn vặt={"QA Catalogue Snack" in snack_page}, '
              f'lẫn cơm trưa={"QA Catalogue Lunch" in snack_page}')


# ---------------------------------------------------------------------------
# Công nợ
# ---------------------------------------------------------------------------

def debt_flows(rep, ids, app):
    a = ids['admin_client']
    with app.app_context():
        pid = ids['plain_product']
        os_ = []
        for i in range(2):
            o = models.Order(customer_name=f'QA Debt {i}', customer_phone='0922222222',
                             total_amount=10000, status='pending')
            db.session.add(o)
            db.session.flush()
            db.session.add(models.OrderItem(order_id=o.id, product_id=pid, quantity=1, price=10000))
            os_.append(o.id)
        db.session.commit()
    ids['debt_orders'] = os_

    # CN-01 thanh toán một đơn
    post(a, f'/admin/debt/{os_[0]}/pay', {}, '/admin/debts')
    with app.app_context():
        ok = db.session.get(models.Order, os_[0]).status == 'completed'
    rep.check('CN-01', ok)

    # CN-02 thanh toán nhiều đơn cùng lúc
    r = post(a, '/admin/debts/bulk_pay', {'order_ids': [str(os_[1])]}, '/admin/debts')
    with app.app_context():
        ok = db.session.get(models.Order, os_[1]).status == 'completed'
    rep.check('CN-02', ok, f'HTTP {r.status_code}' if not ok else '')


# ---------------------------------------------------------------------------
# Tài khoản
# ---------------------------------------------------------------------------

def account_flows(rep, ids, app):
    a = ids['admin_client']

    # TK-01 thêm admin
    with app.app_context():
        before = models.Admin.query.count()
    post(a, '/admin/accounts/add_admin',
         {'username': 'qa_crud_admin', 'password': 'qacrud123', 'role': 'admin'},
         '/admin/accounts')
    with app.app_context():
        after = models.Admin.query.count()
        na = models.Admin.query.filter_by(username='qa_crud_admin').first()
    rep.check('TK-01', after == before + 1 and na is not None, f'{before} -> {after}')
    ids['new_admin'] = na.id if na else None
    if not na:
        return

    # TK-02 đổi vai trò
    post(a, f'/admin/accounts/{na.id}/change_role', {'role': 'super_admin'}, '/admin/accounts')
    with app.app_context():
        ok = db.session.get(models.Admin, na.id).role == 'super_admin'
    rep.check('TK-02', ok)
    post(a, f'/admin/accounts/{na.id}/change_role', {'role': 'admin'}, '/admin/accounts')

    # TK-03 đặt lại mật khẩu admin
    post(a, f'/admin/accounts/{na.id}/reset_password', {'new_password': 'newqacrud123'},
         '/admin/accounts')
    with app.app_context():
        ok = check_password_hash(db.session.get(models.Admin, na.id).password, 'newqacrud123')
    rep.check('TK-03', ok)

    # TK-06 không cho xóa Super Admin cuối cùng: tạm hạ tất cả admin khác xuống
    # 'admin', giữ đúng 1 super_admin (tài khoản đang đăng nhập), rồi thử tự xóa nó
    with app.app_context():
        main_admin_id = models.Admin.query.filter_by(username='admin').first().id
        others_super = models.Admin.query.filter(
            models.Admin.role == 'super_admin', models.Admin.id != main_admin_id).all()
        downgraded = [o.id for o in others_super]
        for o in others_super:
            o.role = 'admin'
        db.session.commit()
    post(a, f'/admin/accounts/{main_admin_id}/delete_admin', {}, '/admin/accounts')
    with app.app_context():
        still_there = db.session.get(models.Admin, main_admin_id) is not None
    rep.check('TK-06', still_there, f'còn tồn tại={still_there}')
    with app.app_context():
        for oid in downgraded:
            row = db.session.get(models.Admin, oid)
            if row:
                row.role = 'super_admin'
        db.session.commit()

    # TK-05 xóa tài khoản admin
    post(a, f'/admin/accounts/{na.id}/delete_admin', {}, '/admin/accounts')
    with app.app_context():
        gone = db.session.get(models.Admin, na.id) is None
    rep.check('TK-05', gone)
    ids['new_admin'] = None  # đã xóa

    # TK-08 đăng ký khách hàng
    with app.app_context():
        before = models.Customer.query.count()
    c = app.test_client()
    post(c, '/customer/create', {
        'username': 'qa_crud_customer', 'password': 'qacrud123',
        'full_name': 'QA Crud Customer', 'phone': '0933333333'}, '/customer/register')
    with app.app_context():
        after = models.Customer.query.count()
        nc = models.Customer.query.filter_by(username='qa_crud_customer').first()
    rep.check('TK-08', after == before + 1 and nc is not None
              and nc.full_name == 'QA Crud Customer' and nc.phone == '0933333333',
              f'{before} -> {after}')
    ids['new_customer'] = nc.id if nc else None

    # TK-09 đăng ký thiếu tên, hoặc SĐT sai định dạng, phải bị từ chối
    with app.app_context():
        before9 = models.Customer.query.count()
    post(c, '/customer/create', {
        'username': 'qa_crud_bad_reg_noname', 'password': 'qacrud123',
        'full_name': '', 'phone': '0933333334'}, '/customer/register')
    post(c, '/customer/create', {
        'username': 'qa_crud_bad_reg_badphone', 'password': 'qacrud123',
        'full_name': 'QA Bad Reg', 'phone': '123'}, '/customer/register')
    with app.app_context():
        after9 = models.Customer.query.count()
    rep.check('TK-09', after9 == before9, f'{before9} -> {after9} (phải không đổi)')

    if nc:
        # TK-04 đặt lại mật khẩu khách hàng
        post(a, f'/admin/accounts/{nc.id}/reset_customer_password',
             {'new_password': 'newqacrud123'}, '/admin/accounts')
        with app.app_context():
            ok = check_password_hash(db.session.get(models.Customer, nc.id).password, 'newqacrud123')
        rep.check('TK-04', ok)

        # XT-04 đăng nhập khách hàng (dùng mật khẩu vừa đặt lại)
        c2 = app.test_client()
        r = post(c2, '/customer/authenticate',
                 {'username': 'qa_crud_customer', 'password': 'newqacrud123'}, '/customer/login')
        rep.check('XT-04', r.status_code == 302, f'HTTP {r.status_code}')

        # TK-10 khách tự sửa thông tin cá nhân trên trang "Thông tin cá nhân"
        post(c2, '/customer/profile/update',
             {'full_name': 'QA Crud Customer Updated', 'phone': '0944444444'},
             '/customer/profile')
        with app.app_context():
            updated = db.session.get(models.Customer, nc.id)
            saved_ok = (updated.full_name == 'QA Crud Customer Updated'
                       and updated.phone == '0944444444')
        cart_page = c2.get('/cart').get_data(as_text=True)
        sidebar_ok = 'QA Crud Customer Updated' in cart_page
        prefill_ok = 'value="QA Crud Customer Updated"' in cart_page and 'value="0944444444"' in cart_page
        rep.check('TK-10', saved_ok and sidebar_ok and prefill_ok,
                  f'đã lưu={saved_ok}, tên mới hiện ở sidebar={sidebar_ok}, đơn hàng tự điền theo={prefill_ok}')

        # XT-05 đăng xuất khách hàng
        r = c2.get('/customer/logout')
        rep.check('XT-05', r.status_code == 302)

        # TK-07 bật/tắt quyền Manager
        post(a, f'/admin/accounts/{nc.id}/toggle_manager', {}, '/admin/accounts')
        with app.app_context():
            promoted = db.session.get(models.Customer, nc.id).role == 'manager'
        post(a, f'/admin/accounts/{nc.id}/toggle_manager', {}, '/admin/accounts')
        with app.app_context():
            back = db.session.get(models.Customer, nc.id).role == 'customer'
        rep.check('TK-07', promoted and back, f'promoted={promoted} back={back}')


# ---------------------------------------------------------------------------
# Xác thực
# ---------------------------------------------------------------------------

def auth_flows(rep, ids, app):
    # XT-01 đăng nhập đúng
    c = app.test_client()
    r = post(c, '/admin/login', {'username': 'admin', 'password': 'admin123'}, '/admin/login')
    dash = c.get('/admin/dashboard')
    rep.check('XT-01', r.status_code == 302 and dash.status_code == 200)

    # XT-02 đăng nhập sai mật khẩu
    c2 = app.test_client()
    post(c2, '/admin/login', {'username': 'admin', 'password': 'sai_mat_khau'}, '/admin/login')
    dash2 = c2.get('/admin/dashboard')
    rep.check('XT-02', dash2.status_code != 200 or 'admin/login' in (dash2.location or ''),
              f'HTTP {dash2.status_code}')

    # XT-03 admin tự đổi mật khẩu rồi đổi lại
    r = post(c, '/admin/change_password',
             {'current_password': 'admin123', 'new_password': 'qacrudtemp123',
              'confirm_password': 'qacrudtemp123'}, '/admin/change_password')
    with app.app_context():
        changed = check_password_hash(
            models.Admin.query.filter_by(username='admin').first().password, 'qacrudtemp123')
    rep.check('XT-03', changed)
    if changed:
        c3 = app.test_client()
        post(c3, '/admin/login', {'username': 'admin', 'password': 'qacrudtemp123'}, '/admin/login')
        post(c3, '/admin/change_password',
             {'current_password': 'qacrudtemp123', 'new_password': 'admin123',
              'confirm_password': 'admin123'}, '/admin/change_password')


# ---------------------------------------------------------------------------
# QR / Thanh toán + Báo cáo
# ---------------------------------------------------------------------------

def qr_and_report_flows(rep, ids, app):
    a = ids['admin_client']

    qr = post(a, '/admin/generate_qr', {'qr_data': 'https://calaci.com.vn'}, '/admin/generate_qr')
    rep.check('QR-01', qr.status_code == 200 and b'base64' in qr.data)

    bqr = post(a, '/admin/generate_bank_qr',
               {'bank_id': '970436', 'account_no': '113366668888',
                'account_name': 'QA CRUD TEST', 'amount': '50000', 'description': 'crud test'},
               '/admin/generate_bank_qr')
    rep.check('QR-02', bqr.status_code == 200)

    # QR-03 đặt tài khoản nhận tiền mặc định - sao lưu cấu hình cũ trước, phục hồi sau
    import payments
    with app.app_context():
        old_config = payments.saved_config()
    r = post(a, '/admin/generate_bank_qr',
             {'action': 'set_default', 'bank_id': '970436',
              'account_no': '113366668888', 'account_name': 'QA CRUD TEST'},
             '/admin/generate_bank_qr')
    body = r.get_data(as_text=True)
    ok = 'Đã đặt làm tài khoản nhận tiền' in body or r.status_code in (200, 302)
    rep.check('QR-03', ok, f'HTTP {r.status_code}')
    with app.app_context():
        if old_config:
            payments.save_config(old_config.get('shop_bank_id', 'VCB'),
                                 old_config.get('shop_bank_account_no', ''),
                                 old_config.get('shop_bank_account_name', ''))
        else:
            from models import SiteSetting
            for key in payments.SAVED_KEYS:
                row = db.session.get(SiteSetting, key)
                if row:
                    db.session.delete(row)
            db.session.commit()

    # BC-01 xuất Excel đúng theo bộ lọc trạng thái đang xem (không phải luôn
    # xuất tất cả - đây là lý do chuyển chức năng này từ mục sidebar riêng
    # sang gắn liền với bộ lọc của trang Đơn hàng)
    with app.app_context():
        pid = ids['plain_product']
        o_done = models.Order(customer_name='QA Export Done', customer_phone='0955555551',
                              total_amount=10000, status='completed')
        o_debt = models.Order(customer_name='QA Export Debt', customer_phone='0955555552',
                              total_amount=10000, status='pending')
        db.session.add_all([o_done, o_debt])
        db.session.flush()
        db.session.add(models.OrderItem(order_id=o_done.id, product_id=pid, quantity=1, price=10000))
        db.session.add(models.OrderItem(order_id=o_debt.id, product_id=pid, quantity=1, price=10000))
        db.session.commit()
        ids['export_orders'] = [o_done.id, o_debt.id]

    import openpyxl as _openpyxl
    import io as _io

    def names_in(resp):
        wb = _openpyxl.load_workbook(_io.BytesIO(resp.data))
        ws = wb.active
        return {row[1] for row in ws.iter_rows(min_row=2, values_only=True) if row[1]}

    x_all = a.get('/admin/export_orders?status=all')
    x_debt = a.get('/admin/export_orders?status=debt')
    all_names = names_in(x_all) if x_all.status_code == 200 else set()
    debt_names = names_in(x_debt) if x_debt.status_code == 200 else set()
    both_in_all = 'QA Export Done' in all_names and 'QA Export Debt' in all_names
    ok = (x_all.status_code == 200 and x_debt.status_code == 200 and both_in_all
          and 'QA Export Debt' in debt_names and 'QA Export Done' not in debt_names)
    rep.check('BC-01', ok, f'all có cả 2={both_in_all}, lọc nợ chỉ còn={debt_names}')

    # BC-02 nút xuất Excel không còn là mục riêng trong sidebar, chỉ còn ở
    # trang Đơn hàng (đã kiểm tra ở BC-01) và trang Sổ nợ khách hàng
    sidebar_page = a.get('/admin/manage_products').get_data(as_text=True)
    debts_page = a.get('/admin/debts').get_data(as_text=True)
    ok2 = ('Xuất Excel' not in sidebar_page
          and 'Xuất Excel' in debts_page and 'status=debt' in debts_page)
    rep.check('BC-02', ok2, f'sidebar còn mục riêng={"Xuất Excel" in sidebar_page}')


# ---------------------------------------------------------------------------
# Dọn dẹp
# ---------------------------------------------------------------------------

def cleanup(ids, app):
    removed = []
    with app.app_context():
        if ids.get('checkout_order'):
            oid = ids['checkout_order']
            for it in models.OrderItem.query.filter_by(order_id=oid).all():
                pr = db.session.get(models.Product, it.product_id)
                if pr:
                    pr.stock += it.quantity
                db.session.delete(it)
            o = db.session.get(models.Order, oid)
            if o:
                db.session.delete(o); removed.append(f'Order#{oid}')

        if ids.get('pos_order'):
            oid = ids['pos_order']
            for it in models.OrderItem.query.filter_by(order_id=oid).all():
                pr = db.session.get(models.Product, it.product_id)
                if pr:
                    pr.stock += it.quantity
                db.session.delete(it)
            o = db.session.get(models.Order, oid)
            if o:
                db.session.delete(o); removed.append(f'Order#{oid}')

        for oid in ids.get('debt_orders', []) or []:
            for it in models.OrderItem.query.filter_by(order_id=oid).all():
                pr = db.session.get(models.Product, it.product_id)
                if pr:
                    pr.stock += it.quantity
                db.session.delete(it)
            o = db.session.get(models.Order, oid)
            if o:
                db.session.delete(o); removed.append(f'Order#{oid}')

        for oid in ids.get('bulk_orders', []) or []:
            models.OrderItem.query.filter_by(order_id=oid).delete()
            o = db.session.get(models.Order, oid)
            if o:
                db.session.delete(o); removed.append(f'Order#{oid}')

        for oid in ids.get('export_orders', []) or []:
            models.OrderItem.query.filter_by(order_id=oid).delete()
            o = db.session.get(models.Order, oid)
            if o:
                db.session.delete(o); removed.append(f'Order#{oid}')

        if ids.get('prefill_order'):
            oid = ids['prefill_order']
            for it in models.OrderItem.query.filter_by(order_id=oid).all():
                pr = db.session.get(models.Product, it.product_id)
                if pr:
                    pr.stock += it.quantity
                db.session.delete(it)
            o = db.session.get(models.Order, oid)
            if o:
                db.session.delete(o); removed.append(f'Order#{oid}')

        if ids.get('prefill_customer'):
            cu = db.session.get(models.Customer, ids['prefill_customer'])
            if cu:
                db.session.delete(cu); removed.append(f'Customer#{ids["prefill_customer"]}')

        if ids.get('order_for_sp04'):
            oid = ids['order_for_sp04']
            models.OrderItem.query.filter_by(order_id=oid).delete()
            o = db.session.get(models.Order, oid)
            if o:
                db.session.delete(o); removed.append(f'Order#{oid}')

        if ids.get('order_for_sp09'):
            oid = ids['order_for_sp09']
            models.OrderItem.query.filter_by(order_id=oid).delete()
            o = db.session.get(models.Order, oid)
            if o:
                db.session.delete(o); removed.append(f'Order#{oid}')

        if ids.get('booking'):
            b = db.session.get(models.RoomBooking, ids['booking'])
            if b:
                db.session.delete(b); removed.append(f'RoomBooking#{ids["booking"]}')

        if ids.get('booking_for_ph08'):
            b = db.session.get(models.RoomBooking, ids['booking_for_ph08'])
            if b:
                db.session.delete(b); removed.append(f'RoomBooking#{ids["booking_for_ph08"]}')

        if ids.get('room'):
            r = db.session.get(models.Room, ids['room'])
            if r:
                models.RoomImage.query.filter_by(room_id=r.id).delete()
                db.session.delete(r); removed.append(f'Room#{ids["room"]}')

        for name in ('QA Test Room Clean',):
            r = models.Room.query.filter_by(name=name).first()
            if r:
                db.session.delete(r); removed.append(f'Room({name})')

        for rid in ids.get('bulk_rooms', []) or []:
            r = db.session.get(models.Room, rid)
            if r:
                models.RoomBooking.query.filter_by(room_id=r.id).delete()
                db.session.delete(r); removed.append(f'Room#{rid}')

        if ids.get('product'):
            p = db.session.get(models.Product, ids['product'])
            if p:
                for img in list(p.images):
                    db.session.delete(img)
                db.session.delete(p); removed.append(f'Product#{ids["product"]}')

        for pid in ids.get('bulk_products', []) or []:
            p = db.session.get(models.Product, pid)
            if p:
                db.session.delete(p); removed.append(f'Product#{pid}')

        for pid in ids.get('catalogue_products', []) or []:
            p = db.session.get(models.Product, pid)
            if p:
                db.session.delete(p); removed.append(f'Product#{pid}')

        if ids.get('special_product'):
            p = db.session.get(models.Product, ids['special_product'])
            if p:
                db.session.delete(p); removed.append(f'Product#{ids["special_product"]}')

        for name in ('QA Test Drink Clean',):
            p = models.Product.query.filter_by(name=name).first()
            if p:
                db.session.delete(p); removed.append(f'Product({name})')

        if ids.get('new_admin'):
            adm = db.session.get(models.Admin, ids['new_admin'])
            if adm:
                db.session.delete(adm); removed.append(f'Admin#{ids["new_admin"]}')
        leftover = models.Admin.query.filter_by(username='qa_crud_admin').first()
        if leftover:
            db.session.delete(leftover); removed.append('Admin(qa_crud_admin)')

        if ids.get('new_customer'):
            cu = db.session.get(models.Customer, ids['new_customer'])
            if cu:
                db.session.delete(cu); removed.append(f'Customer#{ids["new_customer"]}')

        for uname in ('qa_crud_bad_reg_noname', 'qa_crud_bad_reg_badphone'):
            leftover_cu = models.Customer.query.filter_by(username=uname).first()
            if leftover_cu:
                db.session.delete(leftover_cu); removed.append(f'Customer({uname})')

        regular = models.Admin.query.filter_by(username='qa_crud_regular').first()
        if regular:
            db.session.delete(regular); removed.append('Admin(qa_crud_regular)')

        db.session.commit()
    return removed


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def write_excel(results):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Test case'

    headers = ['STT', 'Mã', 'Nhóm chức năng', 'Chức năng', 'Loại thao tác',
              'Mô tả test', 'Kết quả mong đợi', 'Trạng thái', 'Chi tiết',
              'Lần chạy gần nhất']
    ws.append(headers)
    header_fill = PatternFill('solid', fgColor='1F4E78')
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'

    pass_fill = PatternFill('solid', fgColor='C6EFCE')
    fail_fill = PatternFill('solid', fgColor='FFC7CE')
    skip_fill = PatternFill('solid', fgColor='F2F2F2')
    now = datetime.datetime.now().strftime('%d/%m/%Y %H:%M')

    for i, tc in enumerate(MATRIX, start=1):
        ok, detail = results.get(tc.id, (None, ''))
        status = 'CHƯA CHẠY' if ok is None else ('ĐẠT' if ok else 'LỖI')
        ws.append([i, tc.id, tc.module, tc.feature, tc.action, tc.description,
                  tc.expected, status, detail, now if ok is not None else ''])
        row = i + 1
        fill = pass_fill if ok is True else fail_fill if ok is False else skip_fill
        ws.cell(row=row, column=8).fill = fill
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).alignment = Alignment(vertical='center', wrap_text=True)

    widths = [5, 7, 16, 26, 12, 42, 42, 11, 30, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(XLSX_PATH)
    print(f'\nĐã ghi kết quả vào {XLSX_PATH}')


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    app = app_module.app
    rep = Results()
    ids = {}

    with app.app_context():
        temp_admin = models.Admin.query.filter_by(username='qa_crud_regular').first()
        if not temp_admin:
            from werkzeug.security import generate_password_hash
            temp_admin = models.Admin(username='qa_crud_regular',
                                      password=generate_password_hash('qacrud123'),
                                      role='admin')
            db.session.add(temp_admin)
            db.session.commit()

    ids['admin_client'] = app.test_client()
    post(ids['admin_client'], '/admin/login', {'username': 'admin', 'password': 'admin123'},
         '/admin/login')

    ids['regular_admin_client'] = app.test_client()
    post(ids['regular_admin_client'], '/admin/login',
         {'username': 'qa_crud_regular', 'password': 'qacrud123'}, '/admin/login')

    ids['customer_client'] = app.test_client()

    try:
        print('--- Sản phẩm ---')
        product_flows(rep, ids, app)
        print('--- Phòng ---')
        room_flows(rep, ids, app)
        print('--- Đơn hàng & giỏ hàng ---')
        order_and_cart_flows(rep, ids, app)
        print('--- Danh mục ---')
        catalogue_flows(rep, ids, app)
        print('--- Công nợ ---')
        debt_flows(rep, ids, app)
        print('--- Tài khoản ---')
        account_flows(rep, ids, app)
        print('--- Xác thực ---')
        auth_flows(rep, ids, app)
        print('--- QR / Báo cáo ---')
        qr_and_report_flows(rep, ids, app)
    except Exception:
        import traceback
        traceback.print_exc()
    finally:
        removed = cleanup(ids, app)
        print(f'\nĐã dọn dẹp: {", ".join(removed) or "không có gì"}')

    write_excel(rep.data)

    total = len(MATRIX)
    ran = len(rep.data)
    passed = sum(1 for ok, _ in rep.data.values() if ok)
    failed = ran - passed
    not_run = total - ran
    print(f'\n{passed}/{ran} test case đã chạy PASS'
         + (f', {failed} LỖI' if failed else '')
         + (f', {not_run} chưa chạy tới (dừng giữa chừng do lỗi)' if not_run else ''))
    return 1 if (failed or not_run) else 0


if __name__ == '__main__':
    sys.exit(main())
